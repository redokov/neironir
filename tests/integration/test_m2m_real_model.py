"""E2E test for the M2M (Bearer) API against the **real** OPF model.

Red stage (TDD): fails until T08–T11 implement the M2M contract
(particularly the JSON download branch, T10).

Like ``test_pipeline_real_model.py`` these tests are skipped unless:

* ``NEIRONIR_RUN_REAL_MODEL_TESTS=1`` is set;
* the ``opf`` binary is on PATH (or ``NEIRONIR_PRIVACY_FILTER_CMD`` set);
* the real test document exists on disk (CI machines don't have it).

Input document: ``ф2.docx`` — a real contract (docx) with PII
(names, emails, dates, addresses).

``ф3.cleaned.md`` is NOT used as input — it is the reference for the
*expected placeholder style* (``<PRIVATE_PERSON1>``, ``<PRIVATE_EMAIL1>``,
…). The assertions below check that the M2M download (both binary and
JSON/base64) contains placeholders of that style and none of the PII
present in the source ``ф2.docx``.
"""

from __future__ import annotations

import base64
import os
import shutil
import time
from collections.abc import Generator
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

# --- Module-level skip guards -----------------------------------------------
OPF_CMD = os.environ.get("NEIRONIR_PRIVACY_FILTER_CMD") or shutil.which("opf")
RUN_FLAG = os.environ.get("NEIRONIR_RUN_REAL_MODEL_TESTS") == "1"

REAL_DOC_DIR = Path(r"C:\MyProjects\MyTasks\summerizer\Договоры")
INPUT_DOCX = REAL_DOC_DIR / "ф2.docx"

pytestmark = [
    pytest.mark.real_model,
    pytest.mark.skipif(
        not RUN_FLAG,
        reason="NEIRONIR_RUN_REAL_MODEL_TESTS is not set — real-model tests are opt-in",
    ),
    pytest.mark.skipif(
        OPF_CMD is None,
        reason="opf CLI not on PATH and NEIRONIR_PRIVACY_FILTER_CMD not set",
    ),
    pytest.mark.skipif(
        not INPUT_DOCX.is_file(),
        reason=f"real test document not found: {INPUT_DOCX}",
    ),
]

# The API key the M2M client uses in this test. Configured via the
# settings override below.
REAL_MODEL_API_KEY = "real-model-m2m-test-key"
AUTH = {"Authorization": f"Bearer {REAL_MODEL_API_KEY}"}

# PII that is present in ф2.docx and must NOT survive into the result.
SOURCE_PII = (
    "korotaev@motor-invest.ru",
    "r.edokov@open-bs.ru",
)

# Placeholder style taken from the cleaned reference (ф3.cleaned.md):
# <PRIVATE_PERSON1>, <PRIVATE_EMAIL1>, <PRIVATE_DATE1>, …
PLACEHOLDER_PREFIX = "<PRIVATE_"


@pytest.fixture(scope="module")
def client() -> Generator[TestClient, None, None]:
    """TestClient wired to the real OPF subprocess, with M2M keys set."""
    import tempfile

    from neironir.api.dependencies import get_privacy, get_settings, get_storage
    from neironir.config import Settings
    from neironir.main import create_app
    from neironir.privacy.client import SubprocessPrivacyFilterClient
    from neironir.storage.local import LocalStorage

    storage_dir = Path(tempfile.mkdtemp(prefix="m2m_real_model_storage_"))
    storage = LocalStorage(storage_dir)

    privacy = SubprocessPrivacyFilterClient(
        opf_cmd=OPF_CMD.split() if OPF_CMD else None,
        timeout_s=float(os.environ.get("NEIRONIR_PRIVACY_FILTER_TIMEOUT", "600")),
    )

    real_settings = Settings().model_copy(
        update={
            "storage_dir": str(storage_dir),
            "privacy_filter_mode": "subprocess",
            "privacy_filter_cmd": OPF_CMD or "",
            "api_keys": REAL_MODEL_API_KEY,
        }
    )

    app = create_app()
    app.dependency_overrides[get_settings] = lambda: real_settings
    app.dependency_overrides[get_storage] = lambda: storage
    app.dependency_overrides[get_privacy] = lambda: privacy
    app.state.settings = real_settings

    with TestClient(app) as test_client:
        yield test_client

    shutil.rmtree(storage_dir, ignore_errors=True)


def _wait_for_completion(test_client: TestClient, job_id: str, max_wait_s: float = 600.0) -> dict:
    """Poll the job endpoint (with the M2M key) until a terminal state.

    The real OPF model takes seconds per document; the docs recommend
    a poll interval >= 2 s, so we honour it here.
    """
    deadline = time.monotonic() + max_wait_s
    while time.monotonic() < deadline:
        r = test_client.get(f"/api/v1/documents/{job_id}", headers=AUTH)
        assert r.status_code == 200, r.text
        body = r.json()
        if body["status"] in {"completed", "failed"}:
            return body
        time.sleep(2.0)
    raise AssertionError(f"job {job_id} did not complete within {max_wait_s}s")


def _docx_text(data: bytes) -> str:
    """Extract the visible text from downloaded docx bytes."""
    from docx import Document

    document = Document(BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


class TestM2MRealModelFlow:
    """Full M2M flow with a Bearer key on the real neural model:
    upload → poll → binary download + JSON (base64) download."""

    def test_upload_with_key_poll_and_download_both_formats(self, client: TestClient) -> None:
        # --- Upload (Bearer-authenticated) --------------------------------
        with open(INPUT_DOCX, "rb") as fh:
            upload = client.post(
                "/api/v1/documents/",
                files={"file": (INPUT_DOCX.name, fh.read(), "application/octet-stream")},
                headers=AUTH,
            )
        assert upload.status_code == 202, upload.text
        job_id = upload.json()["id"]

        # --- Poll until completed ------------------------------------------
        job = _wait_for_completion(client, job_id)
        assert job["status"] == "completed", job

        # --- Binary download ----------------------------------------------
        binary = client.get(f"/api/v1/documents/{job_id}/download", headers=AUTH)
        assert binary.status_code == 200, binary.text
        assert "wordprocessingml.document" in binary.headers["content-type"]
        assert 'filename="' in binary.headers.get("content-disposition", "")
        binary_text = _docx_text(binary.content)

        # --- JSON (base64) download ---------------------------------------
        json_resp = client.get(
            f"/api/v1/documents/{job_id}/download",
            headers={**AUTH, "Accept": "application/json"},
        )
        assert json_resp.status_code == 200, json_resp.text
        body = json_resp.json()
        assert set(body) == {
            "job_id",
            "filename",
            "ext",
            "media_type",
            "size",
            "content_base64",
        }
        assert body["job_id"] == job_id
        assert body["ext"] == "docx"
        assert body["size"] == len(binary.content)
        json_bytes = base64.b64decode(body["content_base64"])
        assert json_bytes == binary.content
        json_text = _docx_text(json_bytes)

        # --- Both results carry placeholders, not PII ----------------------
        for label, text in (("binary", binary_text), ("json", json_text)):
            assert PLACEHOLDER_PREFIX in text, (
                f"{label} download contains no <PRIVATE_*> placeholders — "
                "the real model did not redact anything"
            )
            for pii in SOURCE_PII:
                assert pii not in text, f"{label} download leaks PII: {pii}"
